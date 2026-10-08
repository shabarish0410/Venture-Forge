"""Owner-approved exports of exact accepted artifact versions; no sharing links."""
import csv
import hashlib
import io
import json
from html import escape


def export_hash(payload):
    return hashlib.sha256(json.dumps(payload, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode()).hexdigest()


class ExportFailure(ValueError):
    pass


def blocks(payload):
    yield "Title", "Venture Forge " + payload["venture"]["name"]
    yield "BodyText", "Purpose: " + payload["purpose"]
    yield "BodyText", "Venture revision: " + str(payload["venture"]["revision"])
    yield "BodyText", "Owner-approved selected records. Assumptions, projections and simulations are not verified venture facts."
    for artifact in payload["artifacts"]:
        yield "Heading1", artifact["title"]
        yield "BodyText", "Classification: " + artifact["classification"] + "; artifact: " + artifact["id"]
        for item in artifact["sections"]:
            yield "Heading2", item["title"]
            if item["description"]: yield "BodyText", item["description"]
            if not item["rows"]: yield "BodyText", "No records supplied yet."
            for row in item["rows"]:
                yield "BodyText", " | ".join(f"{key.replace('_', ' ')}: {value if value is not None else 'Unknown'}" for key, value in row.items())
        for gap in artifact["gaps"]: yield "BodyText", "Open question: " + gap
    yield "Heading1", "Source index"
    for source in payload["sources"]:
        yield "BodyText", f"{source['id']}: {source['title']} - {source['locator']} - captured {source['collected_on']}"
    if not payload["sources"]: yield "BodyText", "No source receipts selected. Treat unsourced entries as founder assumptions."


def render_export(payload, format):
    if format == "json": return json.dumps(payload, indent=2, ensure_ascii=False).encode(), "application/json"
    if format == "csv":
        out = io.StringIO(newline="")
        writer = csv.writer(out)
        writer.writerow(["artifact_id", "application", "classification", "section", "record", "field", "value"])
        def cell(value):
            text = "" if value is None else str(value)
            return "'" + text if text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else text
        for artifact in payload["artifacts"]:
            for section in artifact["sections"]:
                for number, row in enumerate(section["rows"], 1):
                    for key, value in row.items():
                        writer.writerow([cell(v) for v in [artifact["id"], artifact["application"], artifact["classification"], section["title"], number, key, value]])
        return out.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8"
    if format == "docx":
        from docx import Document
        from docx.shared import Pt, RGBColor
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = "VENTURE FORGE / OWNER EXPORT"
        doc.styles["Normal"].font.name = "Calibri"
        doc.styles["Normal"].font.size = Pt(10)
        for name in ["Title", "Heading 1", "Heading 2"]:
            doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        for role, line in blocks(payload):
            if role == "Title": doc.add_paragraph(line, style="Title")
            elif role.startswith("Heading"): doc.add_heading(line, int(role[-1]))
            else: doc.add_paragraph(line)
        doc.sections[0].footer.paragraphs[0].text = "Selected, reviewed versions. Check assumptions and source permissions before distribution."
        out = io.BytesIO(); doc.save(out)
        return out.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from pathlib import Path
    import reportlab
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    out = io.BytesIO()
    styles = getSampleStyleSheet()
    if "ForgeVera" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("ForgeVera", str(Path(reportlab.__file__).parent / "fonts" / "Vera.ttf")))
    for name in ["Title", "Heading1", "Heading2", "BodyText"]:
        styles[name].fontName = "ForgeVera"
        styles[name].textColor = "#111111"
    styles["BodyText"].fontSize = 9
    styles["BodyText"].leading = 13
    story = []
    for role, line in blocks(payload):
        # Built-in fonts have limited glyph coverage; keep mathematical labels readable.
        text = line.replace("×", " x ").replace("–", "-").replace("₹", "INR ")
        supported = pdfmetrics.getFont("ForgeVera").face.charToGlyph
        if any(not c.isspace() and ord(c) not in supported for c in text):
            raise ExportFailure("Some characters are unsupported by the MVP PDF font. Use DOCX or JSON to preserve the original text.")
        story += [Paragraph(escape(text).replace("\n", "<br/>"), styles[role])]
        if role == "BodyText": story.append(Spacer(1, 7))
    def footer(canvas, doc):
        canvas.setFont("Helvetica", 8)
        canvas.drawString(40, 26, "Venture Forge | Selected owner export | Page " + str(doc.page))
    SimpleDocTemplate(out, pagesize=A4, leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=44).build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue(), "application/pdf"
